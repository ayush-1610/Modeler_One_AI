"""Writes the two workbooks the client-data e2e flow uploads: a filled client-data template and a raw workbook.

Software test data (made-up numbers): it proves the read path, never a model. Usage: client_workbooks.py <out_dir>
"""

import io
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from modeler_intake.client_template import build_template

out = Path(sys.argv[1])
wb = load_workbook(io.BytesIO(build_template()))
header = [c.value.rstrip("*") for c in wb["Studies"][1]]
study = {"study_id": "CL-01", "design": "SD", "population": "healthy", "n": 12, "route": "oral", "dose": 10,
         "dose_unit": "mg", "formulation": "ir_tablet", "food_state": "fasted", "lloq": 0.5, "lloq_unit": "ng/ml"}
wb["Studies"].append([study.get(h) for h in header])
for t, v in ((0.5, 20.0), (1, 41.0), (2, 30.5), (4, 12.0), (8, "BLQ")):
    wb["PK_Summary"].append(["CL-01", None, t, "h", "arithmetic_mean", v, None, None, 12, "ng/ml"])
wb["PK_Summary"].append(["CL-01", None, 12, "h", "arithmetic_mean", "1,2", None, None, 12, "ng/ml"])  # refused: decimal comma
wb.save(out / "client-template-filled.xlsx")

raw = Workbook()
sheet = raw.active
sheet.title = "Diss pH 6.8"
for row in (["Time (min)", "Vessel 1", "Vessel 2", "% dissolved mean"], [15, 61, 63, 62], [30, 88, 90, 89]):
    sheet.append(row)
raw.create_sheet("Notes").append(["Shipped with the courier on 1 September"])
raw.save(out / "client-raw.xlsx")
