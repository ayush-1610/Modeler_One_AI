"""T-47: the P3 client-data endpoints — template download and read, triage (person and A4), recipe mapping, gate."""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from modeler_api import project_api
from modeler_api.auth import get_verifier
from modeler_api.main import app
from modeler_project import ArtifactKind, FileProjectStore, Workspace
from modeler_project.brief import empty_brief
from modeler_project.requirements import derive

pytestmark = pytest.mark.req("T-47")
H = {"Authorization": "Bearer t"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class FakeVerifier:
    def verify(self, token):
        return {"sub": "u1", "name": "Dev User", "tenant_id": "t1", "realm_access": {"roles": ["modeler-curator"]},
                "projects": ["*"], "acr": "loa2", "auth_time": int(time.time())}


@pytest.fixture
def setup(tmp_path, monkeypatch):
    store = FileProjectStore(tmp_path)
    monkeypatch.delenv("MODELER_LLM_PROVIDER", raising=False)
    app.dependency_overrides[get_verifier] = lambda: FakeVerifier()
    app.dependency_overrides[project_api.get_project_store] = lambda: store
    ws = Workspace(store, "t1", "p1")
    brief = ws.commit(ArtifactKind.BRIEF, "main", empty_brief("Exampleamide", by="u1").to_content(), actor="u1", reason="start")
    ws.commit(ArtifactKind.REQUIREMENTS, "main", derive(empty_brief("Exampleamide", by="u1")).to_content(),
              derived_from=[brief.ref], actor="system", reason="derived")
    yield TestClient(app), ws
    app.dependency_overrides.clear()


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_template_is_downloaded_filled_and_read_into_client_datasets(setup):
    c, _ws = setup
    r = c.get("/api/v1/client-data/template.xlsx", headers=H)
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(r.content))
    header = [x.value.rstrip("*") for x in wb["Studies"][1]]
    values = {"study_id": "CL-1", "design": "SD", "population": "healthy", "n": 12, "route": "oral", "dose": 10,
              "dose_unit": "mg", "formulation": "ir_tablet", "food_state": "fasted"}
    wb["Studies"].append([values.get(h) for h in header])
    for t, v in ((0.5, 20.0), (1, 41.0), (2, 30.5), (4, 12.0)):
        wb["PK_Summary"].append(["CL-1", None, t, "h", "arithmetic_mean", v, None, None, 12, "ng/ml"])
    up = c.post("/api/v1/projects/p1/client-data", headers=H,
                files=[("files", ("client.xlsx", _save(wb), XLSX)), ("files", ("report.md", b"# Study CL-1\nLLOQ 0.5 ng/ml", "text/markdown"))])
    assert up.status_code == 201, up.text
    view = up.json()["data"]
    files = {f["file"]: f for f in view["files"]}
    workbook, report = files["client.xlsx"], files["report.md"]
    assert workbook["template"] and len(workbook["datasets"]) == 1 and not workbook["issues"]
    assert report["kind"] == "markdown" and not report["datasets"]
    evidence = c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]
    ds = next(d for d in evidence["datasets"] if d["id"] == workbook["datasets"][0])
    assert ds["origin"] == "CLIENT" and ds["provider"] == "CLIENT" and ds["series"][0]["values"] == [20.0, 41.0, 30.5, 12.0]
    # every client item of the default plan (nothing promised yet) is listed; the client's study is unpromised
    assert any("CL-1" in u for u in view["reconciliation"]["unpromised"])


def _raw_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in (["Study ABC-1, 25 mg oral tablet, fasted"], ["Subject", "Time (h)", "Conc (ng/mL)"],
                ["001", 0.5, 10.1], ["001", 1, 22.4], ["001", 2, 15.0], ["002", 0.5, 9.8], ["002", 1, 20.0], ["002", 2, 14.1]):
        ws.append(row)
    notes = wb.create_sheet("Misc")
    notes.append(["Shipped with the courier on 1 September"])
    return _save(wb)


def test_any_other_workbook_is_triaged_classified_and_mapped_by_a_confirmed_recipe(setup):
    c, _ws = setup
    view = c.post("/api/v1/projects/p1/client-data", headers=H, files=[("files", ("raw.xlsx", _raw_workbook(), XLSX))]).json()["data"]
    sub = view["files"][0]
    assert not sub["template"]
    triage = {t["sheet"]: t for t in sub["triage"]}
    assert triage["Sheet1"]["category"] == "PK_INDIVIDUAL" and triage["Sheet1"]["evidence_quote"] == "Subject"
    assert triage["Misc"]["category"] == "OTHER"
    refused = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:triage", headers=H)
    assert refused.status_code == 409 and "agents are off" in refused.json()["detail"]
    view = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Misc:classify", headers=H,
                  json={"category": "PRODUCT_INFO", "reason": "batch list, per the client's email"}).json()["data"]
    assert {t["sheet"]: t["by"] for t in view["files"][0]["triage"]}["Misc"] == "u1"

    proposal = {"tables": [{
        "record_type": "concentration_time", "sheet": "Sheet1", "header_rows": 2, "first_data_row": 3,
        "columns": [{"column": "A", "role": "subject_id"}, {"column": "B", "role": "time"}, {"column": "C", "role": "value"}],
        "time_unit": "h", "value_unit": "ng/mL", "statistic": "individual",
        "constants": [{"key": "study_id", "value": "ABC-1"}, {"key": "analyte", "value": "parent"},
                      {"key": "matrix", "value": "plasma"}, {"key": "dose", "value": "25"}, {"key": "dose_unit", "value": "mg"},
                      {"key": "route", "value": "oral"}, {"key": "formulation", "value": "ir_tablet"},
                      {"key": "food_state", "value": "fasted"}],
        "evidence": [{"cell": "Sheet1!C2", "quote": "ng/mL", "supports": "value unit"},
                     {"cell": "Sheet1!A1", "quote": "25 mg", "supports": "dose"}]}]}
    preview = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                     json={"proposal": proposal, "study": {"n": 2, "design": "SD"}}).json()["data"]
    assert preview["concentrations"] == 6 and preview["studies"] == ["ABC-1"] and preview["ready"], preview
    done = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                  json={"proposal": proposal, "study": {"n": 2, "design": "SD", "purpose": "external_validation"}, "confirm": True})
    assert done.status_code == 200, done.text
    ids = done.json()["data"]["datasets"]
    ds = next(d for d in c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]["datasets"] if d["id"] == ids[0])
    assert ds["origin"] == "CLIENT" and [s["name"] for s in ds["series"]] == ["001", "002"] and ds["study"]["dose_mg"] == 25
    # the purpose the person chose goes on the dataset; the study record stays a valid MS-01 record
    assert ds["purpose"] == "external_validation" and "purpose" not in ds["study"]
    assert ds["quote"].startswith("[Sheet1 row 3]")


def test_a4_classifies_only_with_a_header_quote_code_can_find(setup):
    from modeler_agents.llm import ChatResult, ToolCall
    from modeler_api.client_api import run_triage_job

    c, ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H, files=[("files", ("raw.xlsx", _raw_workbook(), XLSX))]).json()["data"]["files"][0]

    class Scripted:
        provider, model = "test", "s"

        def __init__(self):
            self.turn = 0

        def complete(self, messages, tools):
            self.turn += 1
            calls = () if self.turn > 1 else (
                ToolCall(id="1", name="classify_sheet", arguments={"sheet": "Misc", "category": "PRODUCT_INFO",
                                                                   "cell": "Misc!A1", "quote": "Shipped with the courier"}),
                ToolCall(id="2", name="classify_sheet", arguments={"sheet": "Sheet1", "category": "DISSOLUTION",
                                                                   "cell": "Sheet1!A1", "quote": "dissolution"}))
            return ChatResult(content="" if calls else "done", tool_calls=calls, finish_reason="", usage={}, model="s")

    result = run_triage_job(ws.store, "t1", "p1", sub["id"], model=Scripted())
    # only Misc was undecided by code; the claim about Sheet1 is refused (not offered, and no such quote)
    assert result["classified"] == 1
    triage = {t["sheet"]: t for t in c.get("/api/v1/projects/p1/client-data", headers=H).json()["data"]["files"][0]["triage"]}
    assert triage["Misc"]["category"] == "PRODUCT_INFO" and triage["Misc"]["by"].startswith("agent:")
    assert triage["Sheet1"]["category"] == "PK_INDIVIDUAL"


@pytest.mark.req("T-48")
def test_dissolution_profiles_are_shown_and_a_fit_is_proposed_as_a_release_model(setup):
    import numpy as np

    from pbpk_domain.dissolution import weibull_fraction

    c, _ws = setup
    wb = load_workbook(io.BytesIO(c.get("/api/v1/client-data/template.xlsx", headers=H).content))
    times = (5, 10, 15, 20, 30, 45, 60)
    for t, f in zip(times, weibull_fraction(np.array(times), 25.0, 1.1), strict=True):
        wb["Dissolution"].append(["Tab 10", "TEST", 10, "B7", "USP 2 paddle", 50, "FaSSIF", 6.5, 500, 37, None, t, "min",
                                  *[round(100 * float(f) + d, 2) for d in np.linspace(-1, 1, 12)]])
    view = c.post("/api/v1/projects/p1/client-data", headers=H, files=[("files", ("diss.xlsx", _save(wb), XLSX))]).json()["data"]
    profile = view["dissolution"]["profiles"][0]
    assert profile["release_model"] == "Weibull" and profile["fit"]["t50_min"] == pytest.approx(25.0, rel=0.01)
    r = c.post(f"/api/v1/projects/p1/dissolution/{profile['id']}:propose", headers=H, json={"formulation": "Tab10"})
    assert r.status_code == 201 and len(r.json()["data"]["evidence"]) == 4
    targets = {e["target"] for e in c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]["evidence"]}
    assert {"form.Tab10.type", "form.Tab10.weibull.t50", "form.Tab10.weibull.shape", "form.Tab10.weibull.lag"} <= targets
    bad = c.post(f"/api/v1/projects/p1/dissolution/{profile['id']}:propose", headers=H, json={"formulation": "Tab 1.0"})
    assert bad.status_code == 422 and "no dots" in bad.json()["detail"]


def _be_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reference"
    ws.append(["Study 230-23: plasma desvenlafaxine (ng/mL), 50 mg, fasting"])
    ws.append(["Subject No.", "Period", "Pre-dose", "1.00", "2.00", "4.00 h", "8.00", "24.00"])
    ws.append(["001", "I", "BLQ", 40.5, 80.2, 95.0, 60.1, 12.3])
    ws.append(["002", "II", "BLQ", 35.0, 70.4, 88.8, 55.5, "NS"])
    ws.append(["Mean", None, None, 37.8, 75.3, 91.9, 57.8, None])
    ws.append([])
    ws.append(["BLQ: below 1.0 ng/mL"])
    return _save(wb)


def test_a_sheet_is_shown_with_a_suggested_form_and_read_from_the_form(setup):
    c, _ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H,
                 files=[("files", ("230-23 Fasting Reference.xlsx", _be_workbook(), XLSX))]).json()["data"]["files"][0]
    sheet = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]
    assert sheet["rows"][1][:3] == ["Subject No.", "Period", "Pre-dose"] and sheet["read_before"] == []
    form = sheet["form"]
    assert form["layout"] == "times_across" and form["subject_column"] == "A" and form["last_data_row"] == 4
    assert form["constants"]["study_id"] == "230-23-REF" and form["lloq"] == 1.0
    assert any("'NS'" in n for n in sheet["notes"])

    # "NS" is not a number: the preview names the cell, and nothing is kept
    assert form["mean_row"] == 5
    form["mean_n"] = 2                        # the Mean row under the subjects is read as the study's mean profile
    study = {"n": 2, "design": "SD", "crossover": True, "purpose": "external_validation"}
    preview = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H, json={"form": form, "study": study}).json()["data"]
    assert not preview["ready"] and [i["location"] for i in preview["issues"]] == ["Reference!H4"]
    form["missing_tokens"] = ["NS"]
    preview = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H, json={"form": form, "study": study}).json()["data"]
    assert preview["ready"] and preview["concentrations"] == 11 + 4 and preview["recipe"]["tables"][0]["time_row"] == 2
    sample = preview["sample"]
    assert sample["series"] == ["001 / I", "002 / II", "Mean"] and sample["times"] == [0, 1, 2, 4, 8, 24] and sample["below_lloq"] == 2
    assert sample["rows"][1] == {"series": "001 / I", "time": 1.0, "value": 40.5, "blq": False, "cell": "Reference!D3"}
    done = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                  json={"form": form, "study": study, "confirm": True}).json()["data"]
    ds = next(d for d in c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]["datasets"] if d["id"] == done["datasets"][0])
    assert ds["purpose"] == "external_validation" and ds["study"]["study_id"] == "230-23-REF" and ds["study"]["lloq"] == 1.0
    again = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]
    assert again["read_before"][0]["datasets"] == done["datasets"]
    refused = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H, json={"study": study})
    assert refused.status_code == 422


def test_reading_a_sheet_again_replaces_the_earlier_reading(setup):
    c, _ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H,
                 files=[("files", ("230-23_Fasting_Reference.xlsx", _be_workbook(), XLSX))]).json()["data"]["files"][0]
    form = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]["form"]
    assert form["constants"]["study_id"] == "230-23-REF" and form["constants"]["food_state"] == "fasted"
    form["missing_tokens"] = ["NS"]
    form["mean_n"] = 2
    study = {"n": 2, "design": "SD", "purpose": "model_building"}
    first = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                   json={"form": form, "study": study, "confirm": True}).json()["data"]["datasets"]
    # the purpose was wrong: read again, replacing the first reading
    second = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                    json={"form": form, "study": {**study, "purpose": "external_validation"}, "confirm": True, "replace": True})
    assert second.status_code == 200, second.text
    datasets = {d["id"]: d for d in c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]["datasets"]}
    assert datasets[first[0]]["state"] == "REJECTED" and "read again" in datasets[first[0]]["decision_reason"]
    assert datasets[second.json()["data"]["datasets"][0]]["purpose"] == "external_validation"
    view = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]
    assert [r["datasets"] for r in view["read_before"]] == [second.json()["data"]["datasets"]]


def test_a_triage_start_answers_through_its_typed_model(setup, monkeypatch):
    # phase 6c: no test reached the started answer of :triage (the conftest guard checks each answer)
    from modeler_api import agent_jobs, client_api, jobs

    c, ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H,
                 files=[("files", ("raw.xlsx", _raw_workbook(), XLSX))]).json()["data"]["files"][0]

    class Configured:
        provider, model = "test", "scripted"

    calls = []
    monkeypatch.setattr(agent_jobs, "chat_model", lambda: (Configured(), None))
    monkeypatch.setattr(client_api, "run_triage_job", lambda *args, **kwargs: calls.append(args))
    started = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:triage", headers=H)
    assert started.status_code == 202 and started.json()["data"] == {"status": "RUNNING"}
    for _ in range(200):
        if not jobs.running(ws.store, "t1", "p1", "triage"):
            break
        time.sleep(0.01)
    assert calls and calls[0][3] == sub["id"]


@pytest.mark.req("T-21")
def test_the_reading_agent_proposes_each_unread_data_sheet_and_the_page_shows_its_checked_proposal(setup):
    """R3c: A4 corrects code's first filling of each data sheet; code checks it, and the sheet then opens on the agent's
    reading, marked as the agent's, for the person to check and save. Nothing is read into the project by the agent."""
    from modeler_agents.llm import ChatResult, ToolCall
    from modeler_api.client_api import run_reading_job

    c, ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H,
                 files=[("files", ("230-23_Fasting_Reference.xlsx", _be_workbook(), XLSX))]).json()["data"]["files"][0]
    first = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]
    assert first["proposed_by"] == "code" and first["run_id"] is None
    corrected = {**first["form"], "missing_tokens": ["NS"], "mean_n": 2}     # what code's filling left for a person
    guessed = {**corrected, "constants": {**corrected["constants"], "n": "24"}}  # an n the sheet never states

    class Scripted:
        provider, model = "test", "scripted"

        def __init__(self):
            self.turns = [[("propose_reading", {"form": guessed})], [("propose_reading", {"form": corrected})], []]

        def complete(self, messages, tools):
            calls = self.turns.pop(0)
            return ChatResult(content="done", finish_reason="", model="s", usage={"input_tokens": 1, "output_tokens": 1},
                              tool_calls=tuple(ToolCall(id=str(i), name=n, arguments=a) for i, (n, a) in enumerate(calls)))

    result = run_reading_job(ws.store, "t1", "p1", sub["id"], model=Scripted())
    assert [(s["sheet"], s["proposed"], s["refused"]) for s in result["sheets"]] == [("Reference", True, 1)]
    assert c.get("/api/v1/projects/p1/evidence", headers=H).json()["data"]["datasets"] == []    # nothing read yet
    sheet = c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Reference", headers=H).json()["data"]
    assert sheet["proposed_by"] == f"agent:{result['sheets'][0]['run_id']}" and sheet["run_id"] == result["sheets"][0]["run_id"]
    assert sheet["form"]["missing_tokens"] == ["NS"] and "n" not in sheet["form"]["constants"]
    run = c.get(f"/api/v1/projects/p1/agent-runs/{sheet['run_id']}", headers=H).json()["data"]
    assert run["run"]["agent"] == "A4-sheet-reading" and run["run"]["status"] == "COMPLETED"
    assert any(s["kind"] == "tool_result" and "REJECTED" in str(s["content"]) for s in run["steps"])   # the refusal is kept

    # the person saves the agent's reading once; the sheet is then read, and nothing is left for the agent
    study = {"n": 2, "design": "SD", "purpose": "model_building"}
    saved = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:map", headers=H,
                   json={"form": sheet["form"], "study": study, "confirm": True})
    assert saved.status_code == 200, saved.text
    assert len(saved.json()["data"]["datasets"]) == 1
    left = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:read", headers=H)
    assert left.status_code == 409 and "no data sheet is left to read" in left.json()["detail"]


@pytest.mark.req("T-21")
def test_a_reading_start_runs_in_the_background_and_refuses_the_template_and_unknown_sheets(setup, monkeypatch):
    from modeler_api import agent_jobs, client_api, jobs
    from modeler_intake.client_template import build_template

    c, ws = setup
    sub = c.post("/api/v1/projects/p1/client-data", headers=H,
                 files=[("files", ("230-23_Fasting_Reference.xlsx", _be_workbook(), XLSX))]).json()["data"]["files"][0]
    off = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:read", headers=H)
    assert off.status_code == 409 and "agents are off" in off.json()["detail"]

    class Configured:
        provider, model = "test", "scripted"

    calls = []
    monkeypatch.setattr(agent_jobs, "chat_model", lambda: (Configured(), None))
    monkeypatch.setattr(client_api, "run_reading_job", lambda *args, **kwargs: calls.append((args, kwargs)))
    assert c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:read", headers=H, json={"sheets": ["Nope"]}).status_code == 404
    started = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:read", headers=H, json={"sheets": ["Reference"]})
    assert started.status_code == 202 and started.json()["data"] == {"status": "RUNNING"}
    for _ in range(200):
        if not jobs.running(ws.store, "t1", "p1", "reading"):
            break
        time.sleep(0.01)
    assert calls and calls[0][0][3] == sub["id"] and calls[0][1]["sheets"] == ["Reference"]
    template = c.post("/api/v1/projects/p1/client-data", headers=H,
                      files=[("files", ("filled.xlsx", build_template(), XLSX))]).json()["data"]["files"]
    tid = next(f["id"] for f in template if f["file"] == "filled.xlsx")
    refused = c.post(f"/api/v1/projects/p1/client-data/{tid}:read", headers=H)
    assert refused.status_code == 409 and "cell by cell" in refused.json()["detail"]


@pytest.mark.req("T-42")
def test_the_data_request_lists_what_only_the_client_sends_on_start_and_downloads_as_the_workbook(setup):
    # page review 1 (the owner's decision 2b of 2026-10-29): the CRO's day-one data request, from the data plan
    from modeler_project.requirements import RequirementMatrix, client_request

    c, ws = setup
    expected = client_request(RequirementMatrix.from_content(ws.latest(ArtifactKind.REQUIREMENTS, "main").content))
    assert expected, "the core template asks the client for something even before the brief is filled"
    start = c.get("/api/v1/projects/p1/brief", headers=H).json()["data"]["start"]
    assert [line["req_id"] for line in start["client_request"]] == [line.req_id for line in expected]
    r = c.get("/api/v1/projects/p1/client-data/request.xlsx", headers=H)
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == 'attachment; filename="p1-data-request.xlsx"'
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Request" and wb["Request"]["A1"].value.startswith("Data request: Exampleamide (p1), ")
    rows = [[cell.value for cell in row] for row in wb["Request"].iter_rows(min_row=5)]
    assert [row[1] for row in rows] == [line.what for line in expected]


def test_a_withdrawn_client_file_is_shown_but_not_read_into_the_project(setup):
    # the document library, PR 2: a withdrawn file is kept and opened, but not classified, read or mapped until restored
    from modeler_project.documents import DocumentLibrary

    c, ws = setup
    view = c.post("/api/v1/projects/p1/client-data", headers=H, files=[("files", ("raw.xlsx", _raw_workbook(), XLSX))]).json()["data"]
    sub = view["files"][0]
    DocumentLibrary(ws).withdraw(sub["sha256"], by="u1", reason="the client sent a corrected file")
    assert c.get(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Sheet1", headers=H).status_code == 200
    refused = c.post(f"/api/v1/projects/p1/client-data/{sub['id']}/sheets/Misc:classify", headers=H,
                     json={"category": "PRODUCT_INFO", "reason": "batch list"})
    assert refused.status_code == 409 and "the client sent a corrected file" in refused.json()["detail"]
    assert c.post(f"/api/v1/projects/p1/client-data/{sub['id']}:read", headers=H).status_code == 409
