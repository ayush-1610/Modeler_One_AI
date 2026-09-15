"""Spreadsheets as grids of cells that keep their sheet!A1 reference.

Merged cells: every cell in a merged range carries the top-left value and remembers where it came from,
so a header like "Concentration (ng/mL)" spanning three columns applies to all three.
"""

from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

CellValue = str | float | int | bool | datetime | date | None


def cell_ref(sheet: str, row: int, column: int) -> str:
    return f"{sheet}!{get_column_letter(column)}{row}"


@dataclass(frozen=True)
class Cell:
    sheet: str
    row: int
    column: int
    value: CellValue
    merged_from: str | None = None

    @property
    def ref(self) -> str:
        return cell_ref(self.sheet, self.row, self.column)


@dataclass
class SheetGrid:
    name: str
    max_row: int
    max_column: int
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)

    def cell(self, row: int, column: int | str) -> Cell:
        col = column_index_from_string(column) if isinstance(column, str) else column
        return self.cells.get((row, col)) or Cell(self.name, row, col, None)

    def row_is_empty(self, row: int) -> bool:
        return all(self.cell(row, c).value in (None, "") for c in range(1, self.max_column + 1))

    def preview(self, max_rows: int = 40, max_columns: int = 20) -> str:
        """Compact text view with row numbers and column letters, used when asking AI to propose a mapping."""
        columns = range(1, min(self.max_column, max_columns) + 1)
        lines = ["row\t" + "\t".join(get_column_letter(c) for c in columns)]
        for row in range(1, min(self.max_row, max_rows) + 1):
            values = ["" if self.cell(row, c).value is None else str(self.cell(row, c).value) for c in columns]
            lines.append(f"{row}\t" + "\t".join(values))
        return "\n".join(lines)


@dataclass
class WorkbookGrid:
    file_name: str
    sha256: str
    sheets: dict[str, SheetGrid]

    def sheet(self, name: str) -> SheetGrid:
        if name not in self.sheets:
            raise KeyError(f"sheet {name!r} not found; available: {', '.join(self.sheets)}")
        return self.sheets[name]

    def header_fingerprint(self, sheet: str, header_rows: int) -> str:
        """Fingerprint of a sheet's header region (text cells only), used to recognise a known file layout."""
        grid = self.sheet(sheet)
        parts = [sheet]
        for row in range(1, header_rows + 1):
            for column in range(1, grid.max_column + 1):
                value = grid.cell(row, column).value
                if isinstance(value, str) and value.strip():
                    parts.append(f"{row}:{column}:{value.strip().lower()}")
        return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


def _load_xlsx(path: Path, sha256: str) -> WorkbookGrid:
    workbook = load_workbook(path, data_only=True)
    sheets: dict[str, SheetGrid] = {}
    for ws in workbook.worksheets:
        grid = SheetGrid(name=ws.title, max_row=ws.max_row, max_column=ws.max_column)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    grid.cells[(c.row, c.column)] = Cell(ws.title, c.row, c.column, c.value)
        for merged in ws.merged_cells.ranges:
            origin = grid.cells.get((merged.min_row, merged.min_col))
            if origin is None:
                continue
            for r in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    if (r, col) != (merged.min_row, merged.min_col):
                        grid.cells[(r, col)] = Cell(ws.title, r, col, origin.value, merged_from=origin.ref)
        sheets[ws.title] = grid
    return WorkbookGrid(path.name, sha256, sheets)


def _load_csv(path: Path, sha256: str) -> WorkbookGrid:
    name = path.stem
    grid = SheetGrid(name=name, max_row=0, max_column=0)
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for r, row in enumerate(csv.reader(handle), start=1):
            grid.max_row = r
            grid.max_column = max(grid.max_column, len(row))
            for c, value in enumerate(row, start=1):
                if value != "":
                    grid.cells[(r, c)] = Cell(name, r, c, value)
    return WorkbookGrid(path.name, sha256, {name: grid})


def read_workbook(path: Path, sha256: str) -> WorkbookGrid:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _load_xlsx(path, sha256)
    if suffix == ".csv":
        return _load_csv(path, sha256)
    raise ValueError(f"unsupported spreadsheet type {suffix!r}; convert .xls to .xlsx first")


def as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()
