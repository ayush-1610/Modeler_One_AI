"""Deterministic application of a confirmed mapping recipe to a workbook grid."""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils import column_index_from_string

from modeler_intake.grid import SheetGrid, WorkbookGrid, as_text
from modeler_intake.recipe import MappingRecipe, TableMapping
from modeler_intake.records import ConcentrationObservation, DissolutionObservation, SourceRef
from pbpk_domain.issues import Issue

_LESS_THAN = re.compile(r"^<\s*([0-9]*\.?[0-9]+(?:[eE][-+]?[0-9]+)?)$")


@dataclass(frozen=True)
class ParsedValue:
    value: float | None
    below_lloq: bool = False
    lloq: float | None = None


class UnparseableValueError(ValueError):
    pass


def parse_numeric(raw: object, *, below_lloq_tokens: list[str], decimal_comma: bool = False) -> ParsedValue:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return ParsedValue(None)
    if isinstance(raw, bool):
        raise UnparseableValueError(f"boolean {raw!r} is not a number")
    if isinstance(raw, int | float):
        return ParsedValue(float(raw))
    text = as_text(raw)
    if text.upper() in {t.upper() for t in below_lloq_tokens}:
        return ParsedValue(None, below_lloq=True)
    match = _LESS_THAN.match(text.replace(",", ".") if decimal_comma else text)
    if match:
        return ParsedValue(None, below_lloq=True, lloq=float(match.group(1)))
    if decimal_comma:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        raise UnparseableValueError(f"{text!r} contains a comma; set decimal_comma if the file uses decimal commas")
    try:
        return ParsedValue(float(text))
    except ValueError as exc:
        raise UnparseableValueError(f"{text!r} is not a number") from exc


@dataclass
class ApplyResult:
    concentrations: list[ConcentrationObservation] = field(default_factory=list)
    dissolution: list[DissolutionObservation] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)


def _rows(grid: SheetGrid, table: TableMapping, time_column: int) -> range:
    if table.last_data_row is not None:
        return range(table.first_data_row, table.last_data_row + 1)
    last = table.first_data_row - 1
    for row in range(table.first_data_row, grid.max_row + 1):
        if grid.row_is_empty(row):
            break
        last = row
    return range(table.first_data_row, last + 1)


def _const(table: TableMapping, key: str) -> str | float | None:
    return table.constants.get(key)


def apply_recipe(workbook: WorkbookGrid, recipe: MappingRecipe) -> ApplyResult:
    result = ApplyResult()
    for table in recipe.tables:
        grid = workbook.sheet(table.sheet)
        by_role: dict[str, list] = {}
        for mapping in table.columns:
            by_role.setdefault(mapping.role, []).append(mapping)
        time_col = column_index_from_string(by_role["time"][0].column)

        for row in _rows(grid, table, time_col):
            time_cell = grid.cell(row, time_col)
            try:
                time = parse_numeric(time_cell.value, below_lloq_tokens=[], decimal_comma=table.decimal_comma).value
            except UnparseableValueError as exc:
                result.issues.append(Issue("UNPARSEABLE_TIME", time_cell.ref, str(exc)))
                continue
            if time is None:
                continue

            row_labels = {
                role: as_text(grid.cell(row, by_role[role][0].column).value)
                for role in ("subject_id", "group")
                if role in by_role
            }
            for value_mapping in by_role["value"]:
                value_cell = grid.cell(row, value_mapping.column)
                try:
                    parsed = parse_numeric(
                        value_cell.value, below_lloq_tokens=table.below_lloq_tokens, decimal_comma=table.decimal_comma
                    )
                except UnparseableValueError as exc:
                    result.issues.append(Issue("UNPARSEABLE_VALUE", value_cell.ref, str(exc)))
                    continue
                if parsed.value is None and not parsed.below_lloq:
                    continue

                series = value_mapping.series_label or row_labels.get("subject_id") or row_labels.get("group") or "all"
                cells = {"time": time_cell.ref, "value": value_cell.ref}
                source = SourceRef(file_sha256=workbook.sha256, recipe_id=recipe.recipe_id, recipe_version=recipe.version, cells=cells)

                if table.record_type == "concentration_time":
                    sd = n = None
                    if "sd" in by_role:
                        sd_cell = grid.cell(row, by_role["sd"][0].column)
                        sd = parse_numeric(sd_cell.value, below_lloq_tokens=[], decimal_comma=table.decimal_comma).value
                        cells["sd"] = sd_cell.ref
                    if "n" in by_role:
                        n_cell = grid.cell(row, by_role["n"][0].column)
                        n_value = parse_numeric(n_cell.value, below_lloq_tokens=[]).value
                        n = int(n_value) if n_value is not None else None
                        cells["n"] = n_cell.ref
                    elif _const(table, "n") is not None:
                        n = int(float(_const(table, "n")))
                    dose = _const(table, "dose")
                    result.concentrations.append(
                        ConcentrationObservation(
                            study_id=str(_const(table, "study_id") or ""),
                            analyte=str(_const(table, "analyte") or ""),
                            matrix=str(_const(table, "matrix") or ""),
                            series=series,
                            statistic=table.statistic,
                            time=time,
                            time_unit=table.time_unit,
                            value=parsed.value,
                            unit=table.value_unit,
                            below_lloq=parsed.below_lloq,
                            lloq=parsed.lloq or table.lloq,
                            sd=sd,
                            n=n,
                            dose=float(dose) if dose is not None else None,
                            dose_unit=_const(table, "dose_unit"),
                            route=_const(table, "route"),
                            formulation=_const(table, "formulation"),
                            food_state=_const(table, "food_state"),
                            source=source,
                        )
                    )
                else:
                    if parsed.value is None:
                        result.issues.append(Issue("MISSING_DISSOLUTION_VALUE", value_cell.ref, "dissolution cells cannot be below LLOQ"))
                        continue
                    ph, rpm = _const(table, "ph"), _const(table, "rpm")
                    result.dissolution.append(
                        DissolutionObservation(
                            batch=str(_const(table, "batch") or ""),
                            medium=str(_const(table, "medium") or ""),
                            ph=float(ph) if ph is not None else None,
                            apparatus=_const(table, "apparatus"),
                            rpm=float(rpm) if rpm is not None else None,
                            vessel=series,
                            time=time,
                            time_unit=table.time_unit,
                            percent_dissolved=parsed.value,
                            source=source,
                        )
                    )
    return result
