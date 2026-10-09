"""Text of uploaded documents, page by page, so every extracted value can cite a page and a verbatim quote (plan §6).

Technical proposals, annexes, papers and client files arrive as PDF, Word, Markdown, plain text, CSV or Excel
(owner's answer to §3 #15). Each becomes a list of pages:

- PDF: one page per PDF page (pypdf's text layer). A PDF with no text layer is reported, not guessed: it needs OCR.
- Word (.docx): body paragraphs and tables in document order, cut into pages of about 3 500 characters at paragraph
  boundaries (Word files carry no page breaks in their XML until rendered).
- Markdown / text: split at form feeds, else into ~3 500-character pages at line boundaries.
- CSV / Excel: one page per sheet (cut when long); each row is written as ``[Sheet row 3] A3=… | B3=…`` so a quote of
  a row's text also names its cells.

The bytes themselves are kept unchanged in the vault; this is only the searchable, quotable text.
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from pathlib import PurePath

PAGE_CHARS = 3500
MAX_BYTES = 60 * 1024 * 1024

KINDS = {
    ".pdf": "pdf", ".docx": "docx", ".md": "markdown", ".markdown": "markdown", ".txt": "text", ".text": "text",
    ".csv": "csv", ".tsv": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".xls": "xls", ".png": "image", ".jpg": "image", ".jpeg": "image",
}
MEDIA_TYPES = {
    "pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "markdown": "text/markdown", "text": "text/plain", "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "image": "image/png",
}


_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # the compound-file signature every Excel 97–2003 workbook starts with


class DocumentError(ValueError):
    """The file cannot be read as a document (unsupported type, corrupt, too large)."""


@dataclass(frozen=True)
class DocumentPage:
    page: int     # 1-based
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    sha256: str
    name: str
    kind: str
    media_type: str
    size_bytes: int
    pages: tuple[DocumentPage, ...]
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def page_text(self, page: int) -> str | None:
        return self.pages[page - 1].text if 0 < page <= len(self.pages) else None


def detect_kind(data: bytes, filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".doc":
        raise DocumentError("legacy .doc files are not read; save the file as .docx")
    kind = KINDS.get(suffix)
    if kind is None:
        if data.startswith(b"%PDF"):
            return "pdf"
        raise DocumentError(f"unsupported file type {suffix or '(none)'}: use PDF, DOCX, Markdown, text, CSV, XLSX, PNG or JPEG")
    if kind == "pdf" and not data.startswith(b"%PDF"):
        raise DocumentError(f"{filename} is not a PDF file")
    if kind in ("docx", "xlsx") and not data.startswith(b"PK"):
        raise DocumentError(f"{filename} is not a valid {kind.upper()} file")
    if kind == "xls" and not data.startswith(_OLE2):
        raise DocumentError(f"{filename} is not a valid XLS (Excel 97–2003) file")
    if kind == "image" and not data.startswith((b"\x89PNG", b"\xff\xd8")):
        raise DocumentError(f"{filename} is not a PNG or JPEG image")
    return kind


def _chunk(blocks: list[str], limit: int = PAGE_CHARS) -> list[str]:
    """Join text blocks into pages of at most ~`limit` characters, never splitting a block unless it alone is longer."""
    pages: list[str] = []
    current: list[str] = []
    size = 0
    for block in blocks:
        while len(block) > limit:
            if current:
                pages.append("\n".join(current))
                current, size = [], 0
            pages.append(block[:limit])
            block = block[limit:]
        if size + len(block) > limit and current:
            pages.append("\n".join(current))
            current, size = [], 0
        current.append(block)
        size += len(block) + 1
    if current:
        pages.append("\n".join(current))
    return pages or [""]


def _pdf_pages(data: bytes) -> tuple[list[str], list[str]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        texts = [(page.extract_text() or "") for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentError(f"the PDF could not be read: {exc}") from exc
    warnings = []
    if texts and not any(t.strip() for t in texts):
        warnings.append("the PDF has no text layer (a scan?): OCR is needed before values can be cited from it")
    return texts, warnings


def _docx_blocks(data: bytes) -> list[str]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for a bad file
        raise DocumentError(f"the Word file could not be read: {exc}") from exc
    blocks: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text
            if text.strip():
                blocks.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    return blocks


def _text_pages(data: bytes) -> list[str]:
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    if "\f" in text:
        return [p for p in text.split("\f")] or [""]
    return _chunk(text.split("\n"))


def _cell_ref(row: int, column: int) -> str:
    letters = ""
    while column:
        column, rem = divmod(column - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def _rows_to_blocks(sheet: str, rows) -> list[str]:
    blocks = []
    for r, row in enumerate(rows, start=1):
        cells = [f"{_cell_ref(r, c)}={v}" for c, v in enumerate(row, start=1) if v not in (None, "")]
        if cells:
            blocks.append(f"[{sheet} row {r}] " + " | ".join(cells))
    return blocks


def _sheet_pages(data: bytes, kind: str, name: str) -> list[str]:
    if kind == "csv":
        delimiter = "\t" if name.lower().endswith(".tsv") else ","
        text = data.decode("utf-8-sig", errors="replace")
        return _chunk(_rows_to_blocks(PurePath(name).stem, csv.reader(io.StringIO(text), delimiter=delimiter)))
    if kind == "xls":
        from modeler_intake.grid import read_workbook_bytes

        try:
            grid = read_workbook_bytes(data, name)
        except ValueError as exc:
            raise DocumentError(str(exc)) from exc
        pages = []
        for sheet in grid.sheets.values():
            rows = [[sheet.cells[(r, c)].value if (r, c) in sheet.cells and not sheet.cells[(r, c)].merged_from else None
                     for c in range(1, sheet.max_column + 1)] for r in range(1, sheet.max_row + 1)]
            pages.extend(_chunk(_rows_to_blocks(sheet.name, rows)))
        return pages or [""]
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zip/xml errors of several types
        raise DocumentError(f"the workbook could not be read: {exc}") from exc
    pages: list[str] = []
    for sheet in workbook.worksheets:
        pages.extend(_chunk(_rows_to_blocks(sheet.title, sheet.iter_rows(values_only=True))))
    return pages or [""]


def extract_document(data: bytes, filename: str) -> ExtractedDocument:
    """Read `data` (named `filename`) into quotable pages."""
    if not data:
        raise DocumentError("the file is empty")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"the file is larger than {MAX_BYTES // (1024 * 1024)} MB")
    kind = detect_kind(data, filename)
    warnings: list[str] = []
    if kind == "pdf":
        texts, warnings = _pdf_pages(data)
    elif kind == "docx":
        texts = _chunk(_docx_blocks(data))
    elif kind in ("markdown", "text"):
        texts = _text_pages(data)
    elif kind == "image":
        texts = [""]
        warnings.append("an image has no text: it can be digitized (figure), not quoted")
    else:
        texts = _sheet_pages(data, kind, filename)
    media_type = "image/jpeg" if kind == "image" and data.startswith(b"\xff\xd8") else MEDIA_TYPES[kind]
    return ExtractedDocument(
        sha256=hashlib.sha256(data).hexdigest(), name=PurePath(filename).name, kind=kind,
        media_type=media_type, size_bytes=len(data),
        pages=tuple(DocumentPage(page=i, text=t) for i, t in enumerate(texts, start=1)), warnings=tuple(warnings),
    )
