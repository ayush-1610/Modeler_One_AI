"""The client-data workbook (plan 2026-09-25 §10.2, D-11): a file read with no AI and no ambiguity.

One sheet per kind of data, one header row, one record per row from row 2. The column list below is the single source
of truth: the downloadable template is built from it and an uploaded workbook is read against it, so the two cannot
drift. Reading is deterministic. Every value keeps the cell it came from, a value that does not parse is reported with
its cell and never guessed, and a missing required value is an issue, not a default.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Literal

from openpyxl.utils import get_column_letter

from modeler_intake.apply import UnparseableValueError, parse_numeric
from modeler_intake.grid import WorkbookGrid, as_text
from modeler_intake.recipe import DEFAULT_BELOW_LLOQ_TOKENS
from pbpk_domain.issues import Issue

TEMPLATE_ID = "ModelerOne_ClientData_v1"
ColumnKind = Literal["text", "number", "enum", "concentration"]


@dataclass(frozen=True)
class Column:
    name: str
    kind: ColumnKind = "text"
    required: bool = False
    options: tuple[str, ...] = ()
    help: str = ""


def _c(name: str, kind: ColumnKind = "text", required: bool = False, options: tuple[str, ...] = (), help: str = "") -> Column:
    return Column(name, kind, required, options, help)


YES_NO = ("yes", "no")
ROLES = ("TEST", "RLD", "REFERENCE", "SOLUTION", "OTHER")
ROUTES = ("oral", "iv_bolus", "iv_infusion")
FORMULATIONS = ("solution", "suspension", "ir_tablet", "ir_capsule", "mr", "other")
STATISTICS = ("arithmetic_mean", "geometric_mean", "median")
VARIABILITY = ("SD", "CV%", "SE")
PK_PARAMETERS = ("AUC_last", "AUC_inf", "Cmax", "tmax", "t_half", "Ctrough", "CL/F", "Vz/F")
USES = ("model_building", "internal_validation", "external_validation", "application_verification", "supportive")
TIME_UNITS = ("min", "h", "day")

SHEETS: dict[str, tuple[Column, ...]] = {
    "Studies": (
        _c("study_id", required=True, help="author-year-route-dose or the client's study code; one row per study arm"),
        _c("treatment", help="the arm within a study (e.g. TEST, RLD, fed); empty for a one-arm study"),
        _c("reference", help="protocol / report number"),
        _c("design", "enum", True, ("SD", "MD")),
        _c("crossover", "enum", options=YES_NO),
        _c("population", "enum", True, ("healthy", "patient")),
        _c("special_population", "enum", options=("pediatric", "hepatic_impairment", "renal_impairment", "pregnancy", "elderly")),
        _c("n", "number", True, help="subjects in this arm"),
        _c("age_mean", "number"), _c("age_min", "number"), _c("age_max", "number"),
        _c("weight_mean_kg", "number"), _c("percent_female", "number"),
        _c("ethnicity"), _c("genotype"), _c("co_medication"),
        _c("route", "enum", True, ROUTES),
        _c("dose", "number", True),
        _c("dose_unit", "enum", True, ("mg", "mg/kg")),
        _c("salt_or_base", "enum", options=("base", "salt")),
        _c("infusion_time_min", "number"),
        _c("product", help="as named on the Product sheet"),
        _c("product_role", "enum", options=ROLES),
        _c("formulation", "enum", True, FORMULATIONS),
        _c("strength_mg", "number"), _c("batch"),
        _c("food_state", "enum", True, ("fasted", "fed")),
        _c("meal_type"), _c("meal_kcal", "number"), _c("meal_fat_g", "number"),
        _c("dosing_interval_h", "number"), _c("n_doses", "number"),
        _c("analyte", help="parent or the metabolite's name"),
        _c("matrix", "enum", options=("plasma", "blood", "serum")),
        _c("lloq", "number"), _c("lloq_unit"),
        _c("intended_use", "enum", options=USES, help="what the client intends the study for, if anything"),
    ),
    "PK_Individual": (
        _c("study_id", required=True), _c("treatment"), _c("subject_id", required=True), _c("period"),
        _c("nominal_time", "number", True), _c("actual_time", "number"), _c("time_unit", "enum", True, TIME_UNITS),
        _c("concentration", "concentration", help="a number, or BLQ / <0.5 below the limit of quantification"),
        _c("unit", required=True), _c("blq_flag", "enum", options=YES_NO),
    ),
    "PK_Summary": (
        _c("study_id", required=True), _c("treatment"), _c("time", "number", True), _c("time_unit", "enum", True, TIME_UNITS),
        _c("statistic", "enum", True, STATISTICS), _c("value", "concentration", True), _c("variability", "number"),
        _c("variability_kind", "enum", options=VARIABILITY), _c("n", "number"), _c("unit", required=True),
    ),
    "PK_Parameters": (
        _c("study_id", required=True), _c("treatment"), _c("parameter", "enum", True, PK_PARAMETERS),
        _c("statistic", "enum", options=STATISTICS), _c("value", "number", True), _c("variability", "number"),
        _c("variability_kind", "enum", options=VARIABILITY), _c("unit", required=True, help="AUC as concentration*time, e.g. ng/ml*h"),
    ),
    "Dissolution": (
        _c("product", required=True), _c("role", "enum", True, ROLES), _c("strength_mg", "number"), _c("batch"),
        _c("apparatus", help="USP 1 basket / USP 2 paddle / ..."), _c("rpm", "number"), _c("medium", required=True),
        _c("ph", "number"), _c("volume_ml", "number"), _c("temperature_c", "number"), _c("surfactant"),
        _c("time", "number", True), _c("time_unit", "enum", True, TIME_UNITS),
        *(_c(f"vessel_{i}", "number") for i in range(1, 13)),
        _c("mean", "number"), _c("sd", "number"), _c("n", "number"),
    ),
    "Product": (
        _c("product", required=True), _c("role", "enum", True, ROLES), _c("dosage_form"), _c("strength_mg", "number"),
        _c("release", "enum", options=("IR", "DR", "ER")), _c("particle_size_d10_um", "number"),
        _c("particle_size_d50_um", "number"), _c("particle_size_d90_um", "number"), _c("particle_size_method"),
        _c("polymorph"), _c("salt_form"),
    ),
    "Physchem_InVitro": (
        _c("parameter", required=True, help="the parameter id (phys.logp, bind.fu, …) or its name on the data plan"),
        _c("value", "number", True), _c("unit"), _c("conditions", help="pH, temperature, species, buffer …"),
        _c("method", help="shake flask, equilibrium dialysis, Caco-2 …; measured or predicted"), _c("report_reference"),
    ),
    "Urine_Feces": (
        _c("study_id", required=True), _c("treatment"), _c("matrix", "enum", True, ("urine", "feces")),
        _c("interval_start_h", "number", True), _c("interval_end_h", "number", True), _c("amount", "number", True),
        _c("unit", required=True, help="mg, µmol, or % of dose"), _c("analyte"),
    ),
}

README = (
    TEMPLATE_ID,
    "",
    "Fill one sheet per kind of data you hold; leave the others empty. Row 1 is the header: do not rename, reorder or",
    "delete columns. One record per row from row 2. Empty rows end nothing: every non-empty row is read.",
    "",
    "Columns marked * are required for each row you fill. Enumerated columns accept only the values listed on the",
    "'Allowed values' rows below; anything else is reported with its cell, never guessed.",
    "Numbers: plain numbers with a decimal point (no thousands separators, no units in the cell).",
    "Below the limit of quantification: write BLQ (or <LLOQ, <0.5) in the concentration cell.",
    "Times: the time after the (first) dose, in the unit of the time_unit column.",
    "Studies with several arms (crossover TEST / RLD, fasted / fed): one Studies row per arm, the arm named in",
    "'treatment', and the same study_id + treatment on the PK rows of that arm.",
    "Physchem_InVitro: name the parameter by its id (phys.logp, phys.pka, phys.solubility.ref, bind.fu, dist.bp_ratio,",
    "perm.intestinal, elim.hepatic.<enzyme>.clint …) or by its name in the data plan, and state the conditions.",
    "Not sure where something goes? Upload what you have: other files (PDF, Word, CSV, Markdown) are accepted too.",
)


def build_template() -> bytes:
    """The downloadable client-data workbook: README, then one sheet per kind of data with its header and checks."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    for line in README:
        readme.append([line])
    readme.append([])
    readme.append(["Sheet", "Column", "Required", "Allowed values", "Notes"])
    for sheet, columns in SHEETS.items():
        for col in columns:
            readme.append([sheet, col.name, "yes" if col.required else "", ", ".join(col.options), col.help])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 20
    readme.column_dimensions["B"].width = 24
    readme.column_dimensions["D"].width = 50
    for sheet, columns in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([f"{c.name}*" if c.required else c.name for c in columns])
        ws.freeze_panes = "A2"
        for i, col in enumerate(columns, start=1):
            ws.cell(1, i).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(col.name) + 3)
            if col.options:
                dv = DataValidation(type="list", formula1='"' + ",".join(col.options) + '"', allow_blank=True)
                dv.error = f"{col.name}: one of {', '.join(col.options)}"
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass(frozen=True)
class Row:
    """One read row: typed values by column name, and the cell each came from."""

    sheet: str
    row: int
    values: dict[str, Any]
    cells: dict[str, str]
    below_lloq: frozenset[str] = frozenset()   # columns whose cell said BLQ / <LLOQ (value None)

    @property
    def locator(self) -> str:
        return f"{self.sheet}!A{self.row}:{get_column_letter(max(1, len(SHEETS[self.sheet])))}{self.row}"


@dataclass
class ClientWorkbook:
    file_name: str
    sha256: str
    template: bool
    rows: dict[str, list[Row]] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)
    other_sheets: list[str] = field(default_factory=list)

    def sheet(self, name: str) -> list[Row]:
        return self.rows.get(name, [])


def _header(grid, column: int) -> str:
    return as_text(grid.cell(1, column).value).rstrip("*").strip().lower()


def is_template(workbook: WorkbookGrid) -> bool:
    """A workbook built from the template: its README says so, or a template sheet carries the template's header."""
    readme = workbook.sheets.get("README")
    if readme is not None and as_text(readme.cell(1, 1).value) == TEMPLATE_ID:
        return True
    for name, columns in SHEETS.items():
        grid = workbook.sheets.get(name)
        if grid is not None and [_header(grid, i) for i in range(1, len(columns) + 1)] == [c.name for c in columns]:
            return True
    return False


def _parse(column: Column, raw: Any, ref: str, issues: list[Issue]) -> tuple[Any, bool]:
    """(value, below_lloq). Text is stripped; numbers parse strictly; enums match case-insensitively."""
    if raw is None or as_text(raw) == "":
        return None, False
    if column.kind in ("number", "concentration"):
        try:
            parsed = parse_numeric(raw, below_lloq_tokens=list(DEFAULT_BELOW_LLOQ_TOKENS) if column.kind == "concentration" else [])
        except UnparseableValueError:
            issues.append(Issue("UNPARSEABLE_VALUE", ref, f"{column.name}: {as_text(raw)!r} is not a plain number "
                                                          "(the template takes a decimal point, no units or separators)"))
            return None, False
        return parsed.value, parsed.below_lloq
    text = as_text(raw)
    if column.kind == "enum":
        match = next((o for o in column.options if o.lower() == text.lower()), None)
        if match is None:
            issues.append(Issue("NOT_ALLOWED", ref, f"{column.name}: {text!r} is not one of {', '.join(column.options)}"))
            return None, False
        return match, False
    return text, False


def read_client_workbook(workbook: WorkbookGrid) -> ClientWorkbook:
    """Read every template sheet of `workbook` row by row (header in row 1, data from row 2)."""
    result = ClientWorkbook(file_name=workbook.file_name, sha256=workbook.sha256, template=is_template(workbook))
    for name, grid in workbook.sheets.items():
        columns = SHEETS.get(name)
        if columns is None:
            if name != "README":
                result.other_sheets.append(name)
            continue
        found = {_header(grid, i): i for i in range(1, grid.max_column + 1) if _header(grid, i)}
        missing = [c.name for c in columns if c.name not in found]
        if missing and len(missing) == len(columns):
            result.issues.append(Issue("NO_HEADER", f"{name}!A1", "the header row is not the template's; the sheet is not read"))
            continue
        for col in missing:
            result.issues.append(Issue("MISSING_COLUMN", f"{name}!1:1", f"column {col!r} is missing from the header"))
        unknown = sorted(set(found) - {c.name for c in columns})
        for col in unknown:
            result.issues.append(Issue("UNKNOWN_COLUMN", f"{name}!{get_column_letter(found[col])}1",
                                       f"column {col!r} is not part of {TEMPLATE_ID}; it is not read"))
        rows: list[Row] = []
        for r in range(2, grid.max_row + 1):
            if grid.row_is_empty(r):
                continue
            values: dict[str, Any] = {}
            cells: dict[str, str] = {}
            blq: set[str] = set()
            for col in columns:
                index = found.get(col.name)
                if index is None:
                    continue
                cell = grid.cell(r, index)
                value, below = _parse(col, cell.value, cell.ref, result.issues)
                if value is not None or below:
                    values[col.name], cells[col.name] = value, cell.ref
                if below:
                    blq.add(col.name)
                if col.required and as_text(cell.value) == "":
                    result.issues.append(Issue("MISSING_VALUE", cell.ref, f"{col.name} is required"))
            rows.append(Row(sheet=name, row=r, values=values, cells=cells, below_lloq=frozenset(blq)))
        result.rows[name] = rows
    return result
