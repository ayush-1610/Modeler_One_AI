"""T-41: technical proposals and client files become quotable pages."""

from __future__ import annotations

import io

import pytest

from modeler_intake.documents import DocumentError, extract_document

pytestmark = pytest.mark.req("T-41")


def _pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one text line per page (built by hand: no PDF writer dependency)."""
    objects: list[bytes] = []
    kids = []
    n = len(pages)
    # 1 catalog, 2 pages, 3 font, then (page, content) pairs
    for i, text in enumerate(pages):
        page_obj, content_obj = 4 + 2 * i, 5 + 2 * i
        kids.append(f"{page_obj} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"{page_obj} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_obj} 0 R >> endobj\n".encode())
        objects.append(f"{content_obj} 0 obj << /Length {len(stream)} >> stream\n".encode() + stream + b"\nendstream endobj\n")
    head = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{' '.join(kids)}] /Count {n} >> endobj\n".encode(),
        b"3 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for obj in head + objects:
        offsets.append(len(body))
        body += obj
    xref = len(body)
    body += f"xref\n0 {len(offsets) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        body += f"{off:010d} 00000 n \n".encode()
    body += f"trailer << /Size {len(offsets) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return body


def test_pdf_pages_keep_their_numbers():
    doc = extract_document(_pdf(["Single oral dose of 50 mg", "Client provides dissolution data"]), "proposal.pdf")
    assert doc.kind == "pdf" and len(doc.pages) == 2
    assert "50 mg" in doc.page_text(1) and "dissolution" in doc.page_text(2)
    assert doc.page_text(3) is None


def test_docx_reads_paragraphs_and_tables_in_order():
    import docx

    d = docx.Document()
    d.add_paragraph("Objective: predict food effect for the 10 mg tablet.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Data", "Provider"
    table.cell(1, 0).text, table.cell(1, 1).text = "Dissolution", "Client"
    buf = io.BytesIO()
    d.save(buf)
    doc = extract_document(buf.getvalue(), "Proposal.DOCX")
    assert doc.kind == "docx"
    assert doc.pages[0].text.splitlines() == ["Objective: predict food effect for the 10 mg tablet.", "Data | Provider",
                                              "Dissolution | Client"]


def test_markdown_and_csv_and_xlsx_become_pages():
    md = extract_document(b"# Plan\n\nDose 50 mg\fPage two", "notes.md")
    assert [p.text for p in md.pages] == ["# Plan\n\nDose 50 mg", "Page two"]

    csv_doc = extract_document(b"time,conc\n0.5,12.1\n", "pk.csv")
    assert csv_doc.pages[0].text == "[pk row 1] A1=time | B1=conc\n[pk row 2] A2=0.5 | B2=12.1"

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Dissolution"
    wb.active.append(["Time (min)", "% dissolved"])
    wb.active.append([15, 62.5])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx = extract_document(buf.getvalue(), "client.xlsx")
    assert xlsx.pages[0].text.endswith("[Dissolution row 2] A2=15 | B2=62.5")


def test_long_text_is_cut_at_line_boundaries():
    lines = [f"line {i} " + "x" * 90 for i in range(100)]
    doc = extract_document("\n".join(lines).encode(), "long.txt")
    assert len(doc.pages) > 1
    assert all(len(p.text) <= 3500 for p in doc.pages)
    assert doc.pages[1].text.startswith("line ")


def test_unreadable_files_are_refused_with_a_reason():
    with pytest.raises(DocumentError, match="xlsx"):
        extract_document(b"\xd0\xcf\x11\xe0", "old.xls")
    with pytest.raises(DocumentError, match="not a PDF"):
        extract_document(b"hello", "fake.pdf")
    with pytest.raises(DocumentError, match="not a PNG"):
        extract_document(b"x", "image.png")
    with pytest.raises(DocumentError, match="unsupported"):
        extract_document(b"x", "image.gif")
    figure = extract_document(b"\x89PNG\r\n\x1a\n....", "figure2.png")
    assert figure.kind == "image" and figure.pages[0].text == "" and "digitized" in figure.warnings[0]
    with pytest.raises(DocumentError, match="empty"):
        extract_document(b"", "a.txt")
