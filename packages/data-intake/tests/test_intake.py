from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from modeler_intake.apply import UnparseableValueError, apply_recipe, parse_numeric
from modeler_intake.grid import read_workbook
from modeler_intake.recipe import ColumnMapping, MappingRecipe, TableMapping
from modeler_intake.validate import validate_concentrations, validate_dissolution
from modeler_intake.vault import LocalRawVault


def messy_clinical_workbook(path: Path) -> Path:
    """Wide layout, title row, merged header, units in the header, BLQ text, a blank row, and a footnote."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PK_SAD"
    ws["A1"] = "Study XYZ-101: plasma concentrations after 10 mg oral"
    ws["A3"] = "Time (h)"
    ws["B3"] = "Concentration (ng/mL)"
    ws.merge_cells("B3:D3")
    ws["B4"], ws["C4"], ws["D4"] = "S01", "S02", "S03"
    rows = [(0, "BLQ", "BLQ", "BLQ"), (0.5, 12.1, 9.8, 15.2), (1, 25.4, 22.0, 30.1), (4, 11.0, 8.7, "<0.5"), (24, 0.9, "BLQ", 1.1)]
    for i, row in enumerate(rows, start=5):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    ws["A12"] = "BLQ = below the lower limit of quantification (0.5 ng/mL)"
    wb.create_sheet("Dissolution")
    ds = wb["Dissolution"]
    ds.append(["Batch B-17, USP II, 50 rpm, pH 6.8"])
    ds.append(["Time (min)", "V1", "V2", "Mean"])
    for t, v1, v2 in [(5, 21.0, 25.0), (15, 60.5, 58.0), (30, 88.0, 91.5), (60, 99.0, 101.2)]:
        ds.append([t, v1, v2, (v1 + v2) / 2])
    wb.save(path)
    return path


RECIPE = MappingRecipe(
    recipe_id="xyz-101-sad",
    tables=[
        TableMapping(
            record_type="concentration_time",
            sheet="PK_SAD",
            header_rows=4,
            first_data_row=5,
            columns=[
                ColumnMapping(column="A", role="time"),
                ColumnMapping(column="B", role="value", series_label="S01"),
                ColumnMapping(column="C", role="value", series_label="S02"),
                ColumnMapping(column="D", role="value", series_label="S03"),
            ],
            time_unit="h",
            value_unit="ng/mL",
            lloq=0.5,
            constants={"study_id": "XYZ-101", "analyte": "Example-A", "matrix": "plasma", "dose": 10, "dose_unit": "mg", "route": "oral"},
        ),
        TableMapping(
            record_type="dissolution",
            sheet="Dissolution",
            header_rows=2,
            first_data_row=3,
            columns=[
                ColumnMapping(column="A", role="time"),
                ColumnMapping(column="B", role="value", series_label="V1"),
                ColumnMapping(column="C", role="value", series_label="V2"),
            ],
            time_unit="min",
            value_unit="%",
            constants={"batch": "B-17", "medium": "phosphate buffer", "ph": 6.8, "apparatus": "USP II", "rpm": 50},
        ),
    ],
)


def test_vault_is_content_addressed_write_once_and_verifiable(tmp_path):
    source = messy_clinical_workbook(tmp_path / "client.xlsx")
    vault = LocalRawVault(tmp_path / "vault")
    first = vault.store(source, uploaded_by="user:alice")
    second = vault.store(source, uploaded_by="user:bob")
    assert second == first
    assert vault.verify(first.sha256)
    stored = vault.path(first.sha256)
    with pytest.raises(PermissionError):
        stored.open("wb")


def test_messy_workbook_is_read_with_cell_provenance(tmp_path):
    source = messy_clinical_workbook(tmp_path / "client.xlsx")
    stored = LocalRawVault(tmp_path / "vault").store(source, uploaded_by="user:alice")
    grid = read_workbook(source, stored.sha256)
    sheet = grid.sheet("PK_SAD")
    assert sheet.cell(3, "D").value == "Concentration (ng/mL)"
    assert sheet.cell(3, "D").merged_from == "PK_SAD!B3"

    result = apply_recipe(grid, RECIPE)
    assert result.issues == []
    conc = result.concentrations
    assert len(conc) == 15  # stops at the blank row, never reads the footnote
    s03_4h = next(r for r in conc if r.series == "S03" and r.time == 4)
    assert s03_4h.below_lloq and s03_4h.value is None and s03_4h.lloq == 0.5
    assert s03_4h.source.cells["value"] == "PK_SAD!D8"
    assert s03_4h.source.file_sha256 == stored.sha256
    assert validate_concentrations(conc) == []

    assert len(result.dissolution) == 8
    assert validate_dissolution(result.dissolution) == []


def test_same_layout_is_recognised_by_fingerprint(tmp_path):
    a = read_workbook(messy_clinical_workbook(tmp_path / "a.xlsx"), "a" * 64)
    b_path = messy_clinical_workbook(tmp_path / "b.xlsx")
    b = read_workbook(b_path, "b" * 64)
    assert a.header_fingerprint("PK_SAD", 4) == b.header_fingerprint("PK_SAD", 4)


def test_validation_catches_real_problems(tmp_path):
    grid = read_workbook(messy_clinical_workbook(tmp_path / "client.xlsx"), "c" * 64)
    no_lloq = RECIPE.model_copy(update={"tables": [RECIPE.tables[0].model_copy(update={"lloq": None, "value_unit": "ng per mL"})]})
    codes = {i.code for i in validate_concentrations(apply_recipe(grid, no_lloq).concentrations)}
    assert {"LLOQ_MISSING", "UNKNOWN_UNIT"} <= codes


def test_numeric_parsing_rules():
    tokens = ["BLQ"]
    assert parse_numeric(" 12.5 ", below_lloq_tokens=tokens).value == 12.5
    assert parse_numeric("blq", below_lloq_tokens=tokens).below_lloq
    assert parse_numeric("<0.5", below_lloq_tokens=tokens).lloq == 0.5
    assert parse_numeric("1.234,5", below_lloq_tokens=tokens, decimal_comma=True).value == 1234.5
    with pytest.raises(UnparseableValueError, match="decimal_comma"):
        parse_numeric("1,5", below_lloq_tokens=tokens)
    with pytest.raises(UnparseableValueError):
        parse_numeric("n/a?", below_lloq_tokens=tokens)


def test_recipe_requires_one_time_and_a_value_column():
    with pytest.raises(ValueError, match="exactly one time column"):
        TableMapping(
            record_type="dissolution", sheet="S", header_rows=1, first_data_row=2,
            columns=[ColumnMapping(column="B", role="value"), ColumnMapping(column="C", role="value")],
            time_unit="min", value_unit="%",
        )
    confirmed = RECIPE.model_copy(update={"confirmed_by": "user:alice", "confirmed_at": datetime.now(UTC)})
    assert confirmed.confirmed and not RECIPE.confirmed
