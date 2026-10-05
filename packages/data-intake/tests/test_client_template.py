"""T-47: the client-data template is read with no AI; any other workbook is triaged by its headers."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from modeler_intake.client_template import SHEETS, TEMPLATE_ID, build_template, read_client_workbook
from modeler_intake.grid import read_workbook_bytes
from modeler_intake.triage import SheetCategory, quote_in_header, triage

pytestmark = pytest.mark.req("T-47")


def _filled() -> bytes:
    wb = load_workbook(io.BytesIO(build_template()))
    wb["Studies"].append(["s1", "TEST", "CR-001", "SD", "yes", "healthy", None, 12, 30, 20, 45, 72, 50, None, None, None,
                          "oral", 10, "mg", "base", None, "Tab 10", "TEST", "ir_tablet", 10, "B123", "fasted"])
    wb["PK_Summary"].append(["s1", "TEST", 0.5, "h", "arithmetic_mean", 52.1, 11.0, "SD", 12, "ng/ml"])
    wb["PK_Summary"].append(["s1", "TEST", 1, "h", "arithmetic_mean", "BLQ", None, "SD", 12, "ng/ml"])
    wb["PK_Individual"].append(["s1", "TEST", "001", 1, 0.5, 0.52, "h", "<0.5", "ng/ml", "yes"])
    wb["PK_Individual"].append(["s1", "TEST", "002", 1, "1,5", None, "h", 12, "ng/ml"])      # decimal comma: refused
    wb["Physchem_InVitro"].append(["phys.logp", 2.1, None, "pH 7.4, 25 °C", "shake flask, measured", "R-12"])
    wb["Dissolution"].append(["Tab 10", "tablet", 10])                                   # role not allowed, time missing
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_the_template_lists_every_sheet_with_its_header_and_readme():
    wb = load_workbook(io.BytesIO(build_template()))
    assert wb["README"]["A1"].value == TEMPLATE_ID
    for name, columns in SHEETS.items():
        header = [c.value for c in wb[name][1]]
        assert header == [f"{c.name}*" if c.required else c.name for c in columns]


def test_a_filled_template_is_read_with_cell_references_and_problems_named():
    read = read_client_workbook(read_workbook_bytes(_filled(), "client.xlsx"))
    assert read.template and not read.other_sheets
    study = read.sheet("Studies")[0]
    assert study.values["study_id"] == "s1" and study.values["dose"] == 10 and study.values["crossover"] == "yes"
    assert study.cells["dose"] == "Studies!R2" and study.locator.startswith("Studies!A2:")
    summary = read.sheet("PK_Summary")
    assert summary[0].values["value"] == 52.1 and summary[1].values.get("value") is None and "value" in summary[1].below_lloq
    individual = read.sheet("PK_Individual")
    assert "concentration" in individual[0].below_lloq
    assert read.sheet("Physchem_InVitro")[0].values == {"parameter": "phys.logp", "value": 2.1, "conditions": "pH 7.4, 25 °C",
                                                        "method": "shake flask, measured", "report_reference": "R-12"}
    codes = {(i.code, i.location) for i in read.issues}
    assert ("UNPARSEABLE_VALUE", "PK_Individual!E3") in codes          # 1,5 is not guessed
    assert ("NOT_ALLOWED", "Dissolution!B2") in codes                   # "tablet" is not a role
    assert ("MISSING_VALUE", "Dissolution!G2") in codes                 # medium is required
    assert not any(i.location.startswith("Studies") for i in read.issues)


def test_a_template_sheet_with_a_changed_header_is_reported_not_misread():
    wb = load_workbook(io.BytesIO(build_template()))
    wb["PK_Summary"]["F1"] = "conc"
    wb["PK_Summary"].append(["s1", None, 1, "h", "arithmetic_mean", 5, None, None, 12, "ng/ml"])
    buf = io.BytesIO()
    wb.save(buf)
    read = read_client_workbook(read_workbook_bytes(buf.getvalue(), "client.xlsx"))
    codes = {(i.code, i.location) for i in read.issues}
    assert ("MISSING_COLUMN", "PK_Summary!1:1") in codes and ("UNKNOWN_COLUMN", "PK_Summary!F1") in codes
    assert "value" not in read.sheet("PK_Summary")[0].values


def _sheet(wb: Workbook, title: str, rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    for row in rows:
        ws.append(row)


def test_any_other_workbook_is_triaged_by_its_headers_with_the_deciding_cell_quoted():
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Diss pH 6.8", [["Batch B1 in phosphate buffer"], ["Time (min)", "Vessel 1", "Vessel 2", "% dissolved mean"],
                               [15, 61, 63, 62]])
    _sheet(wb, "PK conc", [["Subject", "Time (h)", "Plasma conc (ng/mL)"], ["001", 0.5, 12.1]])
    _sheet(wb, "NCA", [["Treatment", "AUC0-t (ng*h/mL)", "Cmax (ng/mL)", "Tmax (h)"], ["A", 540, 121, 1.0]])
    _sheet(wb, "Notes", [["Shipped on 2026-09-01"]])
    buf = io.BytesIO()
    wb.save(buf)
    workbook = read_workbook_bytes(buf.getvalue(), "client-raw.xlsx")
    result = {t.sheet: t for t in triage(workbook)}
    assert result["Diss pH 6.8"].category is SheetCategory.DISSOLUTION
    assert result["PK conc"].category is SheetCategory.PK_INDIVIDUAL and result["PK conc"].evidence_quote == "Subject"
    assert result["NCA"].category is SheetCategory.PK_PARAMETERS
    assert result["Notes"].category is SheetCategory.OTHER
    # an agent's claim is checked against the sheet's own header cells
    assert quote_in_header(workbook, "NCA", "NCA!C1", "Cmax")
    assert not quote_in_header(workbook, "NCA", "NCA!C1", "AUC") and not quote_in_header(workbook, "NCA", "NCA!C2", "121")


def test_the_template_is_triaged_by_sheet_name_and_its_empty_sheets_are_skipped():
    workbook = read_workbook_bytes(_filled(), "client.xlsx")
    assert {t.sheet: t.category for t in triage(workbook)} == {
        "Studies": SheetCategory.STUDIES, "PK_Individual": SheetCategory.PK_INDIVIDUAL,
        "PK_Summary": SheetCategory.PK_SUMMARY, "Dissolution": SheetCategory.DISSOLUTION,
        "Physchem_InVitro": SheetCategory.PHYSCHEM_INVITRO}
