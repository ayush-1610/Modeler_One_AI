"""T-47: a proposed mapping recipe (from the Data Mapping agent or the sheet form) is applied and checked by code."""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from modeler_intake.grid import read_workbook
from modeler_intake.recipe import CONSTANT_KEYS, NUMERIC_CONSTANTS, TableMapping
from modeler_intake.recipe_review import (
    ColumnProposal,
    ConstantProposal,
    Evidence,
    RecipeProposal,
    TableProposal,
    review_proposal,
)

pytestmark = pytest.mark.req("T-47")


def messy_workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "PK_SAD"
    ws["A1"] = "Study XYZ-101: plasma concentrations after 10 mg oral"
    ws["A3"] = "Time (h)"
    ws["B3"] = "Concentration (ng/mL)"
    ws.merge_cells("B3:C3")
    ws["B4"], ws["C4"] = "S01", "S02"
    for i, row in enumerate([(0, "BLQ", "BLQ"), (1, 25.4, 22.0), (4, 11.0, "<0.5")], start=5):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    ws["A10"] = "BLQ = below the lower limit of quantification (0.5 ng/mL)"
    path = tmp_path / "client.xlsx"
    wb.save(path)
    return read_workbook(path, "d" * 64)


def proposal(**table_overrides) -> RecipeProposal:
    table = dict(
        record_type="concentration_time",
        sheet="PK_SAD",
        header_rows=4,
        first_data_row=5,
        columns=[
            ColumnProposal(column="A", role="time"),
            ColumnProposal(column="B", role="value", series_label="S01"),
            ColumnProposal(column="C", role="value", series_label="S02"),
        ],
        time_unit="h",
        value_unit="ng/mL",
        lloq=0.5,
        below_lloq_tokens=["BLQ"],
        constants=[
            ConstantProposal(key="study_id", value="XYZ-101"),
            ConstantProposal(key="analyte", value="Example-A"),
            ConstantProposal(key="matrix", value="plasma"),
            ConstantProposal(key="dose", value="10"),
            ConstantProposal(key="dose_unit", value="mg"),
        ],
        evidence=[
            Evidence(cell="PK_SAD!B3", quote="ng/mL", supports="value unit"),
            Evidence(cell="PK_SAD!A10", quote="0.5 ng/mL", supports="LLOQ"),
            Evidence(cell="PK_SAD!A1", quote="XYZ-101", supports="study id"),
        ],
    )
    table.update(table_overrides)
    return RecipeProposal(tables=[TableProposal(**table)])


def test_good_proposal_is_ready_for_confirmation(tmp_path):
    review = review_proposal(messy_workbook(tmp_path), proposal(), "xyz-101")
    assert review.issues == [] and review.ready_for_confirmation
    assert len(review.concentrations) == 6
    assert review.recipe.tables[0].constants["dose"] == 10.0
    assert set(review.recipe.fingerprints) == {"PK_SAD"}


def test_unsupported_evidence_and_open_questions_block_confirmation(tmp_path):
    workbook = messy_workbook(tmp_path)
    wrong = proposal(evidence=[Evidence(cell="PK_SAD!A10", quote="1.0 ng/mL", supports="LLOQ")])
    review = review_proposal(workbook, wrong, "xyz-101")
    assert [i.code for i in review.issues] == ["INVALID_EVIDENCE"]

    no_lloq = RecipeProposal(tables=proposal(lloq=None, evidence=[]).tables, questions_for_reviewer=["What is the LLOQ?"])
    review = review_proposal(workbook, no_lloq, "xyz-101")
    assert "LLOQ_MISSING" in {i.code for i in review.issues}
    assert not review.ready_for_confirmation and review.questions == ["What is the LLOQ?"]


def test_blank_constants_are_not_stated_and_a_word_for_a_number_is_named(tmp_path):
    workbook = messy_workbook(tmp_path)
    constants = [ConstantProposal(key="study_id", value="XYZ-101"), ConstantProposal(key="analyte", value="Example-A"),
                 ConstantProposal(key="matrix", value="plasma"), ConstantProposal(key="dose", value=" "),
                 ConstantProposal(key="route", value="")]
    review = review_proposal(workbook, proposal(constants=constants), "xyz-101")
    assert review.ready_for_confirmation and {"dose", "route"}.isdisjoint(review.recipe.tables[0].constants)
    review = review_proposal(workbook, proposal(constants=[*constants, ConstantProposal(key="strength_mg", value="fifty")]), "x")
    assert [i.code for i in review.issues] == ["INVALID_RECIPE"]
    assert "PK_SAD: constant strength_mg must be a number, not 'fifty'" in review.issues[0].message


def test_one_list_of_constant_keys():
    # coupling C9: the proposal, the recipe's schema text and the numeric conversion all read recipe.CONSTANT_KEYS
    assert ConstantProposal.model_json_schema()["properties"]["key"]["enum"] == list(CONSTANT_KEYS)
    assert NUMERIC_CONSTANTS <= set(CONSTANT_KEYS)
    assert TableMapping.model_fields["constants"].description.endswith(", ".join(CONSTANT_KEYS))
