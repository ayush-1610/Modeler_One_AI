"""T-47: client files are read into datasets and evidence that point at their cells, and reconciled with the plan."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from modeler_intake.client_template import build_template
from modeler_project import ArtifactKind, FileProjectStore, Workspace
from modeler_project.brief import Citation, FieldStatus, empty_brief
from modeler_project.brief_ops import agent_set
from modeler_project.client_data import ClientDataError, close_register, ingest, reconcile
from modeler_project.dataset_register import datasets
from modeler_project.documents import DocumentLibrary
from modeler_project.evidence_register import items
from modeler_project.requirements import RequirementOverride, derive

pytestmark = pytest.mark.req("T-47")
CITE = (Citation(doc_sha256="a" * 64, page=1, quote="quoted from the proposal"),)


def _brief():
    b = empty_brief("Exampleamide", by="u")
    for path, value in [
        ("drug.modality", "small_molecule"), ("drug.bcs_class", "III"), ("qoi.applications", ["APP-14"]),
        ("products[0].name", "Test 10 mg tablet"), ("products[0].role", "TEST"), ("products[0].release", "IR"),
        ("scenarios[0].population", "healthy adults"), ("scenarios[0].route", "oral"), ("scenarios[0].dose", 10),
        ("data_plan[0].item", "dissolution in three media"), ("data_plan[0].category", "dissolution"),
        ("data_plan[0].provider", "CLIENT"),
        ("data_plan[1].item", "pivotal BE study individual data"), ("data_plan[1].category", "rld_data"),
        ("data_plan[1].provider", "CLIENT"), ("data_plan[1].purpose", "external_validation"),
        ("data_plan[2].item", "logP and fu measured in house"), ("data_plan[2].category", "physchem"),
        ("data_plan[2].provider", "CLIENT"),
    ]:
        b = agent_set(b, path, value=value, unit=None, citations=CITE, status=FieldStatus.EXTRACTED, confidence="A", by="a1")
    return b


def _workbook() -> bytes:
    wb = load_workbook(io.BytesIO(build_template()))
    studies = wb["Studies"]
    header = [c.value.rstrip("*") for c in studies[1]]

    def study(**values):
        studies.append([values.get(h) for h in header])

    study(study_id="BE-01", treatment="TEST", design="SD", crossover="yes", population="healthy", n=24, route="oral",
          dose=10, dose_unit="mg", formulation="ir_tablet", product="Test 10 mg tablet", product_role="TEST",
          food_state="fasted", lloq=0.5, lloq_unit="ng/ml", intended_use="external_validation")
    study(study_id="BE-01", treatment="RLD", design="SD", crossover="yes", population="healthy", n=24, route="oral",
          dose=25, dose_unit="mg", formulation="ir_tablet", product="Brand 25 mg", product_role="RLD", food_state="fasted",
          intended_use="external_validation")
    for subject, values in (("001", (0, 40.2, 31.0, 12.5)), ("002", ("BLQ", 45.1, 33.3, 10.2))):
        for t, c in zip((0, 1, 2, 6), values, strict=True):
            wb["PK_Individual"].append(["BE-01", "TEST", subject, 1, t, None, "h", c, "ng/ml"])
    for t, mean, sd in ((1, 42.6, 3.5), (2, 32.1, 1.6), (6, 11.3, 1.6)):
        wb["PK_Summary"].append(["BE-01", "RLD", t, "h", "arithmetic_mean", mean, sd, "SD", 24, "ng/ml"])
    wb["PK_Parameters"].append(["BE-01", "RLD", "Cmax", "arithmetic_mean", 42.6, 3.5, "SD", "ng/ml"])
    wb["PK_Summary"].append(["GHOST", None, 1, "h", "arithmetic_mean", 5, None, None, 12, "ng/ml"])   # no Studies row
    wb["Physchem_InVitro"].append(["phys.logp", 1.2, None, "logP octanol/water, pH 7.4", "shake flask, measured", "R-7"])
    wb["Physchem_InVitro"].append(["Fraction unbound in plasma", 0.09, None, "human plasma, 1 µM", "equilibrium dialysis", "R-8"])
    for medium, ph in (("0.1 N HCl", 1.2), ("acetate", 4.5)):
        for t, v in ((15, 55), (30, 85), (45, 97)):
            wb["Dissolution"].append(["Test 10 mg tablet", "TEST", 10, "B1", "USP 2 paddle", 50, medium, ph, 900, 37, None, t,
                                      "min", v, v + 2, v - 1])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _project(tmp_path):
    ws = Workspace(FileProjectStore(tmp_path), "t1", "p1")
    brief = _brief()
    matrix = derive(brief)
    matrix_ref = ws.commit(ArtifactKind.REQUIREMENTS, "main", matrix.to_content(), actor="u", reason="derived").ref
    return ws, DocumentLibrary(ws), brief, matrix, matrix_ref


def test_a_filled_template_becomes_client_datasets_and_evidence_quoted_from_their_rows(tmp_path):
    ws, library, brief, matrix, _ = _project(tmp_path)
    data = _workbook()
    sub = ingest(ws, library, data, "client-2026-10.xlsx", by="u", matrix=matrix, brief=brief)
    assert sub["template"] and len(sub["datasets"]) == 2 and len(sub["evidence"]) == 2
    found = {d.study["study_id"]: d for d in datasets(ws)}
    test, rld = found["BE-01-TEST"], found["BE-01-RLD"]
    assert test.origin.value == "CLIENT" and test.provider == "CLIENT" and test.extraction == "CELL"
    assert [s.name for s in test.series] == ["001", "002"] and test.series[1].values[0] is None   # BLQ kept as such
    assert test.study["lloq"] == 0.5 and test.study["statistic"] == "individual" and test.purpose == "external_validation"
    assert rld.series[0].error_kind == "SD" and rld.reported[0].parameter == "Cmax" and rld.study["dose_mg"] == 25
    assert test.quote.startswith("[PK_Individual row 2]") and "Studies!A2" in test.source.locator
    fu, logp = sorted(items(ws), key=lambda e: e.target)
    assert (fu.target, fu.req_id, fu.value) == ("bind.fu", "REQ-bind.fu", 0.09)
    assert logp.target == "phys.logp" and logp.quote.startswith("[Physchem_InVitro row 2]") and "B2=1.2" in logp.quote
    assert logp.confidence == "A" and logp.provider == "CLIENT" and logp.source.locator == "Physchem_InVitro!A2:F2"
    codes = {i["code"] for i in sub["issues"]}
    assert "UNKNOWN_STUDY" in codes
    assert any("25 mg is not a dose" in m for m in sub["brief_mismatches"])
    assert any("'Brand 25 mg' is not one of the brief's products" in m for m in sub["brief_mismatches"])
    assert len(sub["dissolution"]) == 6
    # the same bytes again: the same submission, nothing proposed twice
    again = ingest(ws, library, data, "client-2026-10.xlsx", by="u", matrix=matrix, brief=brief)
    assert again["repeat"] and len(datasets(ws)) == 2


def test_reconciliation_names_what_was_promised_delivered_partial_or_missing(tmp_path):
    ws, library, brief, matrix, matrix_ref = _project(tmp_path)
    ingest(ws, library, _workbook(), "client.xlsx", by="u", matrix=matrix, brief=brief)
    rows = {r.req_id: r for r in reconcile(ws, matrix).rows}
    release = rows["REQ-form.release@test-10-mg-tablet"]
    assert release.status == "PARTIAL" and release.detail.startswith("2 of 3 promised media")
    assert rows["REQ-vbe.be_study"].status == "DELIVERED"
    assert rows["REQ-phys.logp"].status == "DELIVERED"
    # fu is a binding item the plan leaves to the literature: the client's value is delivered but not promised
    assert any(u.startswith("bind.fu = 0.09") for u in reconcile(ws, matrix).unpromised)
    missing = [r.req_id for r in reconcile(ws, matrix).blocking()]
    assert missing and all(rows[m].status == "MISSING" for m in missing)
    with pytest.raises(ClientDataError, match="still missing"):
        close_register(ws, matrix_ref, matrix, by="u")
    # R-07: the person skips what the client will not send (or asks for a literature cross-check); then P3 closes
    overrides = tuple(RequirementOverride(req_id=m, status="NOT_AVAILABLE", reason="not in the client's scope", by="u")
                      for m in missing)
    matrix = derive(brief, overrides, previous=matrix)
    matrix_ref = ws.commit(ArtifactKind.REQUIREMENTS, "main", matrix.to_content(), actor="u", reason="skipped").ref
    close_register(ws, matrix_ref, matrix, by="u", printed_name="Dr U")
    assert ws.phases()["P3"] == "APPROVED"


def _plain_dissolution_workbook() -> bytes:
    """A client's own layout (not the template): one sheet per product, a header row, time and six vessels."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for sheet, offset in (("Test", 0.0), ("Brand", 2.0)):
        ws = wb.create_sheet(sheet)
        ws.append(["Time (min)", *(f"V{i}" for i in range(1, 7))])
        for t, v in ((15, 40.0), (30, 70.0), (60, 92.0), (120, 99.0)):
            ws.append([t, *(v + offset + d for d in (-1.0, -0.5, 0.0, 0.0, 0.5, 1.0))])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _map_dissolution(ws, sub: dict, data: bytes, sheet: str, **constants) -> None:
    from modeler_intake.apply import apply_recipe
    from modeler_intake.grid import read_workbook_bytes
    from modeler_intake.recipe import ColumnMapping, MappingRecipe, TableMapping
    from modeler_intake.validate import validate_dissolution
    from modeler_project.client_data import record_mapping

    recipe = MappingRecipe(recipe_id=f"{sub['id']}-{sheet}", tables=[TableMapping(
        record_type="dissolution", sheet=sheet, header_rows=1, first_data_row=2, time_unit="min", value_unit="%",
        columns=[ColumnMapping(column="A", role="time"),
                 *(ColumnMapping(column=c, role="value", series_label=f"V{i}") for i, c in enumerate("BCDEFG", start=1))],
        constants={"batch": "B1", "medium": "phosphate", "ph": 6.8, "apparatus": "USP 2 paddle", "rpm": 50, **constants})])
    result = apply_recipe(read_workbook_bytes(data, "dissolution.xlsx"), recipe)
    assert result.issues == [] and validate_dissolution(result.dissolution) == []
    record_mapping(ws, sub["id"], recipe=recipe.model_dump(mode="json"), dataset_ids=[],
                   dissolution=[o.model_dump(mode="json") for o in result.dissolution], by="u")


def test_mapped_dissolution_sheets_count_for_the_release_model_and_the_test_vs_reference_item(tmp_path):
    from modeler_project.dissolution_register import comparisons, profiles

    ws, library, brief, matrix, _ref = _project(tmp_path)
    data = _plain_dissolution_workbook()
    sub = ingest(ws, library, data, "dissolution.xlsx", by="u", matrix=matrix, brief=brief)
    rows = {r.req_id: r for r in reconcile(ws, matrix).rows}
    assert rows["REQ-vbe.rld_dissolution"].status == "MISSING" and rows["REQ-form.release@test-10-mg-tablet"].status == "MISSING"

    # the Test sheet, named the way the brief names the product (case and spacing aside): the release item is served,
    # the test-vs-reference item is partial until a reference profile in the same medium arrives
    _map_dissolution(ws, sub, data, "Test", product="test 10 MG tablet", role="test", strength_mg=10)
    rows = {r.req_id: r for r in reconcile(ws, matrix).rows}
    assert rows["REQ-form.release@test-10-mg-tablet"].delivered == ("phosphate pH 6.8",)
    vbe = rows["REQ-vbe.rld_dissolution"]
    assert vbe.status == "PARTIAL" and "no TEST and RLD / REFERENCE pair" in vbe.detail
    found = {p["key"]["role"]: p for p in profiles(ws)}
    assert found["TEST"]["key"]["product"] == "test 10 MG tablet" and found["TEST"]["key"]["strength_mg"] == 10

    _map_dissolution(ws, sub, data, "Brand", product="Brand 10 mg", role="RLD")
    rows = {r.req_id: r for r in reconcile(ws, matrix).rows}
    assert rows["REQ-vbe.rld_dissolution"].status == "DELIVERED"
    assert rows["REQ-vbe.rld_dissolution"].delivered == ("phosphate pH 6.8",)
    pair = comparisons(ws)["comparisons"]
    assert len(pair) == 1 and pair[0]["condition"] == "phosphate pH 6.8"
    assert not any(u.startswith("dissolution profiles") for u in reconcile(ws, matrix).unpromised)


def test_dissolution_for_a_product_the_brief_does_not_name_says_so(tmp_path):
    ws, library, brief, matrix, _ref = _project(tmp_path)
    data = _plain_dissolution_workbook()
    sub = ingest(ws, library, data, "dissolution.xlsx", by="u", matrix=matrix, brief=brief)
    _map_dissolution(ws, sub, data, "Test", product="ER 50 mg", role="TEST")
    release = {r.req_id: r for r in reconcile(ws, matrix).rows}["REQ-form.release@test-10-mg-tablet"]
    assert release.status == "MISSING" and "'ER 50 mg', not for 'Test 10 mg tablet'" in release.detail


def test_an_item_the_client_cannot_send_is_taken_from_the_literature_and_no_longer_holds_p3(tmp_path):
    ws, library, brief, matrix, matrix_ref = _project(tmp_path)
    ingest(ws, library, _workbook(), "client.xlsx", by="u", matrix=matrix, brief=brief)
    blocking = reconcile(ws, matrix).blocking()
    first = blocking[0]
    assert first.kind and first.target                     # the page says how each item would be delivered
    # a cross-check alone does not close it: it counts once a literature value is accepted, and the refusal says so
    matrix = derive(brief, (RequirementOverride(req_id=first.req_id, cross_check=True, reason="ask the literature", by="u"),),
                    previous=matrix)
    with pytest.raises(ClientDataError, match="counts once a literature value is accepted"):
        close_register(ws, matrix_ref, matrix, by="u")
    # taking it from the literature instead moves it to P2: P3 no longer waits for it
    matrix = derive(brief, (RequirementOverride(req_id=first.req_id, provider="LITERATURE", reason="client has none", by="u"),),
                    previous=matrix)
    assert first.req_id not in [r.req_id for r in reconcile(ws, matrix).blocking()]
    assert matrix.get(first.req_id).provider == "LITERATURE"
